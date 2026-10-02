"""
Read an insurer's own premium calculator (.xlsx) and write what the quoting engine needs.

    python scripts/gb-calculators/build.py            # writes src/lib/gb/calculators/*.json

Run once per insurer per plan year, when the insurer issues a new calculator. Needs openpyxl.

Why this exists. The rate tables were extracted from brochure PDFs by models, and the result
was 17% of QBE's rate points and half of Income's: QBE's hospital cover is priced by plan,
hospital type AND bed count, and only one of its four ward blocks was ever stored; four QBE
products were absent; Income's GST-inclusive rates were having GST added on top; and both
insurers' dependants were labelled as unpriceable when their calculators price them at the
employee rate. The calculator is the insurer's own statement of what a premium is, so it is
the source, read deterministically. No model is involved.

Each output file carries four things:

  rates   every rate point, exactly as the calculator's tables hold it
  plans   the tiers a broker chooses between, with their ward and co-payment attributes
  rules   what the calculator's FORMULAS do to a rate: age basis, GST, eligibility
  cases   expected premiums, computed by re-implementing each calculator formula directly
          against the sheet cells — a separate code path from `rates`, so a test that runs
          the engine over `rates` and compares with `cases` is checking the engine against
          the calculator, not against itself

Broker-chosen variants are written into the plan code ("Plan 1 · Government 4-bedded"), so the
quote wizard, which picks one plan per product, needs no change. Member attributes the broker
does not choose — occupation class — go in `dimensions`, which the engine matches on.
"""
import hashlib, json, os, re, sys
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.normpath(os.path.join(HERE, '..', '..', 'src', 'lib', 'gb', 'calculators'))
DOWNLOADS = os.path.expanduser('~/Downloads')

SEP = ' · '   # plan-code variant separator; the base plan is everything before the first one


def clean(x):
    """Strip floating-point noise without changing a value the calculator can produce:
    495.00000000000006 -> 495.0, 10437.336000000001 -> 10437.336."""
    return round(float(x), 6)


def sha256(path):
    with open(path, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


def bands(floors, last_max):
    """Approximate-MATCH floors -> inclusive [min, max] bands, exactly as MATCH(age, floors, 1)."""
    out = []
    for i, lo in enumerate(floors):
        hi = floors[i + 1] - 1 if i + 1 < len(floors) else last_max
        out.append((lo, hi))
    return out


def edge_ages(floors, last_max):
    """The ages where a pricing mistake would show: each band floor, one either side of it, and
    either side of the oldest age the table prices. Every age in between prices the same as its
    floor, so checking all of them adds size, not coverage."""
    ages = {0, last_max, last_max + 1}
    for f in floors:
        ages.update({f - 1, f, f + 1})
    return sorted(a for a in ages if a >= 0)


def require_text(wb, needle, why, prefer=()):
    """Fail the build if the workbook no longer says what a rule depends on. GST treatment is
    read from the insurer's own words; a new calculator that drops them must stop the build
    rather than carry last year's treatment forward silently."""
    # The current year's sheet first: QBE keeps last year's table (Table_Jan26) in the same file.
    sheets = [wb[n] for n in prefer if n in wb.sheetnames] + [ws for ws in wb.worksheets if ws.title not in prefer]
    for ws in sheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and needle.lower() in c.value.lower():
                    return f'{ws.title}!{c.coordinate}'
    sys.exit(f'STOP: "{needle}" not found in the workbook. {why}')


def match_floor(age, floors):
    """MATCH(age, floors, 1): index of the largest floor <= age, or None below the first."""
    idx = None
    for i, f in enumerate(floors):
        if age >= f:
            idx = i
    return idx


# ════════════════════════════════════════════════════════════════════════════════════════════
# QBE — Steadfast MCare+ calculator v20052026
# ════════════════════════════════════════════════════════════════════════════════════════════

def qbe():
    path = os.path.join(DOWNLOADS, 'Steadfast Calculator v20052026.xlsx')
    wb = load_workbook(path, data_only=True)
    t = wb['Table']
    v = lambda r, c: t.cell(r, c).value
    gst_cell = require_text(wb, 'Rates above exclude GST', 'QBE GST treatment is read from this note.', prefer=('Table',))

    rates, plans, cases = [], [], []
    WARDS = [  # (first column, hospital type as shown, beds)
        (3,  'Private Hospital',                    '1-Bedded'),
        (8,  'Private Hospital',                    '4-Bedded'),
        (13, 'Government/Restructured Hospital',    '1-Bedded'),
        (18, 'Government/Restructured Hospital',    '4-Bedded'),
    ]
    short = {'Private Hospital': 'Private', 'Government/Restructured Hospital': 'Government'}

    # ── Hospital & Surgical. Calculator!N16:
    #    IF(M=0%, INDEX(Table!C8:V20, MATCH(age, A8:A20, 1), col), 0.8 * INDEX(...))
    #    where col = MATCH(plan) + MATCH(hospital) + MATCH(beds) - 2 picks one of 4 ward blocks.
    GHS = 'Group Hospital & Surgical (GHS)'
    floors = [v(r, 1) for r in range(8, 21)]
    gbands = bands(floors, 85)
    renewal_floor = 76          # "* For renewal only. The last entry age for this plan is 75."
    for first_col, hosp, beds in WARDS:
        for p in range(5):
            col = first_col + p
            base = v(4, col)                                  # "Plan 1".."Plan 5"
            for coins, factor in (('0%', 1.0), ('20%', 0.8)):
                code = f'{base}{SEP}{short[hosp]} {beds.lower()}' + (f'{SEP}20% co-insurance' if coins == '20%' else '')
                any_rate = False
                for i, r in enumerate(range(8, 21)):
                    cell = v(r, col)
                    if not isinstance(cell, (int, float)):
                        continue                              # "N/A": Plan 5 has no private ward
                    any_rate = True
                    lo, hi = gbands[i]
                    rates.append(dict(product_code=GHS, plan_code=code, member_type=None,
                                      band_label=str(v(r, 2)).strip(), age_min=lo, age_max=hi,
                                      premium=clean(cell * factor), renewal_only=lo >= renewal_floor,
                                      dimensions={}))
                if any_rate:
                    plans.append(dict(product_code=GHS, plan_code=code, plan_name=code,
                                      hospital_type=short[hosp], beds=beds, co_payment=coins,
                                      annual_limit=v(5, col)))
                    for age in edge_ages(floors, 85):
                        k = match_floor(age, floors)
                        cell = v(8 + k, col) if k is not None and age <= 85 else None
                        exp = clean(cell * factor) if isinstance(cell, (int, float)) else None
                        for rel in ('self', 'spouse', 'child'):
                            cases.append(dict(product_code=GHS, plan_code=code, age=age, relationship=rel,
                                              expected=exp, renewal_only=(k is not None and floors[k] >= renewal_floor)))

    # ── Outpatient GP. Calculator!Q16: INDEX(Table!D54:E57, MATCH(age, A54:A57, 1), MATCH(plan, D52:E52, 0))
    GP = 'Group Outpatient GP (GP)'
    gp_floors = [v(r, 1) for r in range(54, 58)]
    gp_bands = bands(gp_floors, 85)
    for c, copay in ((4, '0'), (5, '10')):                    # Calculator!P16: Plan 1 -> 0, else 10
        code = v(52, c)
        plans.append(dict(product_code=GP, plan_code=code, plan_name=code, hospital_type=None, beds=None,
                          co_payment=f'S${copay} per visit', annual_limit=None))
        for i, r in enumerate(range(54, 58)):
            lo, hi = gp_bands[i]
            rates.append(dict(product_code=GP, plan_code=code, member_type=None, band_label=str(v(r, 2)).strip(),
                              age_min=lo, age_max=hi, premium=clean(v(r, c)), renewal_only=lo >= 76, dimensions={}))
        for age in edge_ages(gp_floors, 85):
            k = match_floor(age, gp_floors)
            exp = clean(v(54 + k, c)) if k is not None and age <= 85 else None
            for rel in ('self', 'spouse', 'child'):
                cases.append(dict(product_code=GP, plan_code=code, age=age, relationship=rel, expected=exp,
                                  renewal_only=(k is not None and gp_floors[k] >= 76)))

    # ── Outpatient Specialist. Calculator!T16: 0% -> Table!D60:F63, 20% -> Table!D66:F69
    SP = 'Group Outpatient Specialist (SP)'
    sp_floors = [v(r, 1) for r in range(60, 64)]
    sp_bands = bands(sp_floors, 85)
    for first_row, coins in ((60, '0%'), (66, '20%')):
        for c in (4, 5, 6):
            base = v(59, c)
            code = base + (f'{SEP}20% co-insurance' if coins == '20%' else '')
            plans.append(dict(product_code=SP, plan_code=code, plan_name=code, hospital_type=None, beds=None,
                              co_payment=coins, annual_limit=None))
            for i in range(4):
                r = first_row + i
                lo, hi = sp_bands[i]
                rates.append(dict(product_code=SP, plan_code=code, member_type=None, band_label=str(v(r, 2)).strip(),
                                  age_min=lo, age_max=hi, premium=clean(v(r, c)), renewal_only=lo >= 76, dimensions={}))
            for age in edge_ages(sp_floors, 85):
                k = match_floor(age, sp_floors)
                exp = clean(v(first_row + k, c)) if k is not None and age <= 85 else None
                for rel in ('self', 'spouse', 'child'):
                    cases.append(dict(product_code=SP, plan_code=code, age=age, relationship=rel, expected=exp,
                                      renewal_only=(k is not None and sp_floors[k] >= 76)))

    # ── Dental. Calculator!W16: Non-Panel -> Table!C94:E97, Panel -> Table!F94:H97
    GD = 'Group Dental (GD)'
    gd_floors = [v(r, 1) for r in range(94, 98)]
    gd_bands = bands(gd_floors, 85)
    for first_col, panel in ((3, 'Non-panel'), (6, 'Panel')):
        for p in range(3):
            c = first_col + p
            code = f'{v(93, c)}{SEP}{panel}'
            plans.append(dict(product_code=GD, plan_code=code, plan_name=code, hospital_type=None, beds=None,
                              co_payment=None, annual_limit=None))
            for i, r in enumerate(range(94, 98)):
                lo, hi = gd_bands[i]
                rates.append(dict(product_code=GD, plan_code=code, member_type=None, band_label=str(v(r, 2)).strip(),
                                  age_min=lo, age_max=hi, premium=clean(v(r, c)), renewal_only=lo >= 76, dimensions={}))
            for age in edge_ages(gd_floors, 85):
                k = match_floor(age, gd_floors)
                exp = clean(v(94 + k, c)) if k is not None and age <= 85 else None
                for rel in ('self', 'spouse', 'child'):
                    cases.append(dict(product_code=GD, plan_code=code, age=age, relationship=rel, expected=exp,
                                      renewal_only=(k is not None and gd_floors[k] >= 76)))

    # ── Personal Accident. Calculator!Z16:
    #    IF(Category="Employee", INDEX(Table!C83:F85, MATCH(class, B83:B85, 0), MATCH(plan, C81:F81, 0)), "N/A")
    #    No age in the formula; employees only.
    GPA = 'Group Personal Accident (GPA)'
    for c in range(3, 7):
        code = v(81, c)
        plans.append(dict(product_code=GPA, plan_code=code, plan_name=f'{code} (S${int(v(82, c)):,} sum insured)',
                          hospital_type=None, beds=None, co_payment=None, annual_limit=v(82, c)))
        for r in range(83, 86):
            cls = re.search(r'\d', str(v(r, 2))).group(0)
            rates.append(dict(product_code=GPA, plan_code=code, member_type='employee', band_label='All ages',
                              age_min=0, age_max=None, premium=clean(v(r, c)), renewal_only=False,
                              dimensions={'occupation_class': cls}))
            for age in (20, 45, 70):
                cases.append(dict(product_code=GPA, plan_code=code, age=age, relationship='self',
                                  occupation_class=cls, expected=clean(v(r, c)), renewal_only=False))
                cases.append(dict(product_code=GPA, plan_code=code, age=age, relationship='spouse',
                                  occupation_class=cls, expected=None, renewal_only=False))

    rules = {
        'age_basis': 'last birthday',
        'gst_treatment': {'treatment': 'exclusive', 'conversion_factor': None, 'read_from': gst_cell},
        'renewal_only_bands': [{'band': [76, 85], 'text': '* For renewal only. The last entry age for this plan is 75.'}],
        'source': 'Steadfast Calculator v20052026.xlsx, read deterministically on 2 Oct 2026',
        'notes': [
            'Rates above exclude GST (Table!B49, B77, B89).',
            '20% co-insurance on hospital and specialist cover is priced at 0.8 x the 0% rate, as Calculator!N16 and the Table notes state.',
            'Personal accident covers employees only (Calculator!Z16 returns N/A for dependants), priced by occupation class 1-3.',
            'Dependants are priced at the same rate as employees for hospital, GP, specialist and dental: no formula distinguishes them.',
            'NOT APPLIED: the note "10% discount applicable for 4-Bedded Private Hospital option" (Table!B48). The calculator prices 4-bedded private from its own column and applies no further discount. Confirm with QBE before relying on either.',
            'NOT APPLIED: the note that Plan 5 is only for 4-bedded government wards. The calculator prices government 1-bedded Plan 5 (S$171 at age 0-25). Confirm with QBE.',
            'Members with chronic disease or major illness are subject to further underwriting (Table!B46).',
        ],
    }
    return dict(insurer='QBE Insurance (Singapore) Pte Ltd', rate_table_id='f07f6289-eb5a-4303-8a7b-882f8aef514c',
                source={'file': os.path.basename(path), 'sha256': sha256(path)},
                rules=rules, rates=rates, plans=plans, cases=cases)


# ════════════════════════════════════════════════════════════════════════════════════════════
# Income — Employees FlexCare calculator (w.e.f. Feb 2026)
# ════════════════════════════════════════════════════════════════════════════════════════════

def income():
    path = os.path.join(DOWNLOADS, 'Income_Calculator - Enhanced Employees FlexCare-NB-Rnwl (w.e.f. Feb2026) - 50 pax.xlsx')
    wb = load_workbook(path, data_only=True)
    rt = wb['Rates']
    v = lambda r, c: rt.cell(r, c).value
    gst_cell = require_text(wb, 'Inclusive of 9% GST', 'Income GST treatment is read from this note.', prefer=('Rates',))

    # Working!M4..S4 = VLOOKUP(age, Rates!A3:AI15, column). Rows 4..15 are the bands; column A the floors.
    floors = [v(r, 1) for r in range(4, 16)]                  # 0,31,36,...,73,76
    ibands = bands(floors, 200)

    # The column for each product and plan, from the calculator's own map (Rates!C17:F36).
    PRODUCTS = [
        ('Group Hospital and Surgical (GHS)',             [(f'Plan {p}', 2 + p, {}) for p in range(1, 8)]),
        ('Group Term Life (GTL)',                         [(f'Plan {p}', 9 + p, {}) for p in range(1, 6)]),
        ('Group Critical Illness (Accelerated) (GCI)',    [(f'Plan {p}', 14 + p, {}) for p in range(1, 6)]),
        ('Group Outpatient Primary Care (GOPC)',          [('Plan 1', 29, {}), ('Plan 2', 30, {})]),
        ('Group Outpatient Specialist Care (GOSC)',       [('Plan 1', 31, {}), ('Plan 2', 32, {})]),
        ('Group Dental (GD)',                             [('Plan 1', 33, {}), ('Plan 2', 34, {})]),
    ]
    # GPA: Working!K4 = VLOOKUP(plan & class, Rates!E18:F26) — "11" is Plan 1, Class 1.
    gpa_map = {}
    for r in range(18, 27):
        key, col = v(r, 5), v(r, 6)
        if key is not None and col is not None:
            gpa_map[str(int(key))] = int(col)
    gpa_cols = [(f'Plan {k[0]}', c, {'occupation_class': k[1]}) for k, c in sorted(gpa_map.items())]
    PRODUCTS.append(('Group Personal Accident (GPA)', gpa_cols))

    # "Eligibility: Up to age 69 last birthday for new applicants. Renewal is up at age 75 last
    #  birthday for GPA, GHS, GP, SP & Dent only." (Import Notes!D32)
    RENEWAL_ONLY = {'Group Hospital and Surgical (GHS)', 'Group Personal Accident (GPA)',
                    'Group Outpatient Primary Care (GOPC)', 'Group Outpatient Specialist Care (GOSC)', 'Group Dental (GD)'}

    rates, plans, cases = [], [], []
    seen_plans = set()
    for product, cols in PRODUCTS:
        for code, col, dims in cols:
            if (product, code) not in seen_plans:
                seen_plans.add((product, code))
                plans.append(dict(product_code=product, plan_code=code, plan_name=code, hospital_type=None,
                                  beds=None, co_payment=None, annual_limit=None))
            for i, r in enumerate(range(4, 16)):
                cell = v(r, col)
                if not isinstance(cell, (int, float)):
                    continue                                  # "Exceed Age"
                lo, hi = ibands[i]
                rates.append(dict(product_code=product, plan_code=code, member_type=None,
                                  band_label=f'{lo} to {hi}' if hi < 200 else f'{lo}+',
                                  age_min=lo, age_max=hi if hi < 200 else None, premium=clean(cell),
                                  renewal_only=(product in RENEWAL_ONLY and lo >= 70), dimensions=dims))
            # Expected, straight from the sheet with Working!C4's eligibility applied.
            for age in sorted(set(edge_ages(floors, 75)) | {15, 16, 24, 25, 69, 70}):
                k = match_floor(age, floors)
                raw = v(4 + k, col) if k is not None else None
                for rel in ('self', 'spouse', 'child'):
                    ineligible = (rel == 'child' and age > 24) or (rel != 'child' and (age > 75 or age < 16))
                    exp = None if ineligible or not isinstance(raw, (int, float)) else clean(raw)
                    case = dict(product_code=product, plan_code=code, age=age, relationship=rel, expected=exp,
                                renewal_only=(product in RENEWAL_ONLY and k is not None and floors[k] >= 70))
                    if dims.get('occupation_class'):
                        case['occupation_class'] = dims['occupation_class']
                    cases.append(case)

    rules = {
        'age_basis': 'last birthday',
        'gst_treatment': {'treatment': 'inclusive', 'conversion_factor': 1.09, 'read_from': gst_cell},
        'renewal_only_bands': [{'band': [70, 75], 'text': 'Up to age 69 last birthday for new applicants. Renewal up to 75 for GPA, GHS, GP, SP & Dent only.'}],
        'eligibility': {'child_max_age': 24, 'adult_min_age': 16, 'adult_max_age': 75},
        'occupation_class_rules': {'excluded_classes': [4]},
        'source': 'Income_Calculator - Enhanced Employees FlexCare-NB-Rnwl (w.e.f. Feb2026) - 50 pax.xlsx, read deterministically on 2 Oct 2026',
        'notes': [
            'Rates are inclusive of 9% GST (Rates!H18, Step 1!Y8, Premium Breakdown!P6). The engine strips it before adding GST, so it is counted once.',
            'Employees, spouses and children are priced from the same rate grid: Working!M4 looks up by age alone.',
            'Not eligible: children over 24, employees or spouses over 75 or under 16 (Working!C4).',
            'Personal accident is priced by plan and occupation class (Working!K4).',
            'Riders: Dental and Primary Care need GHS; Specialist Care needs Primary Care; Critical Illness needs Term Life and takes its plan (Import Notes!D37-D40). Not yet enforced by the quote wizard.',
            'Minimum 2 employees per product (Import Notes!D31). Underwriting above S$150k GTL / S$100k GCI for employees, S$100k / S$50k for dependants, and for ages 65-69 on GTL/GCI (D34-D36).',
            'Plans 6 and 7 are for Work Permit and S Pass holders only, and are the only plans meeting MOM enhanced medical insurance requirements (Import Notes!D26-D28).',
        ],
    }
    return dict(insurer='Income Insurance Limited', rate_table_id='2410c7cd-e03f-4c3c-bb4d-6b62508ae936',
                source={'file': os.path.basename(path), 'sha256': sha256(path)},
                rules=rules, rates=rates, plans=plans, cases=cases)


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    for name, fn in (('qbe-steadfast-2026', qbe), ('income-flexcare-2026', income)):
        data = fn()
        with open(os.path.join(OUT, f'{name}.json'), 'w') as f:
            json.dump(data, f, separators=(',', ':'))
        priced = sum(1 for c in data['cases'] if c['expected'] is not None)
        print(f"{name}: {len(data['rates'])} rates, {len(data['plans'])} plans, "
              f"{len(data['cases'])} cases ({priced} priced) — {data['source']['file']}")
