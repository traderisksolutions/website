"""
Load insurers from TRS EB Calculator V1.2 into the Pricing Matrix.

    python scripts/gb-calculators/v12_load.py            # dry run
    python scripts/gb-calculators/v12_load.py --apply    # create or replace the rate tables

For insurers where V1.2 is the best source on file — RCC (Raffles Corporate Care Enhanced III) and
Singlife MyBenefits Plus as of 2 Oct 2026. QBE and Income are loaded from their own calculators by
build.py/load.py instead, and V1.2 agrees with those on all 578 shared rates.

V1.2 is TRS's own workbook, not an insurer document, so every table loaded here carries a
verification level the Pricing Matrix shows beside its premiums:

  brochure     every rate was found printed in the insurer's brochure (RCC)
  unverified   no insurer document on file to check it against (Singlife)

What V1.2 gets right, checked rather than assumed: net premiums exclude GST, with term life and
critical illness GST-exempt; RCC is priced on age next birthday and the rest on age last birthday;
renewal-only bands are marked; occupation class drives personal accident.

Needs openpyxl, and NEXT_PUBLIC_SUPABASE_URL / SUPABASE_SERVICE_KEY in the environment.
"""
import datetime, hashlib, json, os, re, subprocess, sys
from openpyxl import load_workbook

APPLY = '--apply' in sys.argv
V12 = os.path.expanduser('~/Downloads/TRS EB Calculator V1.2.xlsx')
SB = os.environ['NEXT_PUBLIC_SUPABASE_URL'] + '/rest/v1'
K = os.environ['SUPABASE_SERVICE_KEY']

INSURERS = {
    'RCC Enhanced III': dict(
        company='RAFFLES HEALTH INSURANCE', display='Raffles Health Insurance',
        product_name='Raffles Corporate Care Enhanced III', plan_year=2025, age_basis='next_birthday',
        brochure='RCC Enhanced III Brochure (14 July 2025)_Final.pdf',
        verification={'status': 'brochure',
                      'basis': 'Every rate in TRS EB Calculator V1.2 was found printed in the RCC brochure of 14 Jul 2025 (291 of 291). Position in the table was not checked.'},
        gst_note='RCC brochure p.17: "Prevailing GST applies to all plans, except GTL and GCI (accelerated and additional)." Premium pages are marked "Exclude prevailing GST".',
        age_note='RCC brochure p.17: "Age, if applicable, refers to age next birthday."'),
    'Singlife MyBenefits Plus': dict(
        company='SINGAPORE LIFE LTD.', display='Singapore Life Ltd.',
        product_name='Singlife MyBenefits Plus', plan_year=2022, age_basis='last_birthday',
        brochure=None,
        verification={'status': 'unverified',
                      'basis': 'Only TRS EB Calculator V1.2 (citing brochure COMP/2022/06/MKT/551, pp.9-10). No Singlife document is on file to check against.'},
        gst_note='As recorded in V1.2: net of GST; term life and critical illness GST-exempt. Not confirmed against a Singlife document.',
        age_note='As recorded in V1.2: age last birthday. Not confirmed against a Singlife document.'),
}

# V1.2 cover names -> canonical products (mirrors resolveProduct; pinned in gb-resolve.test.ts).
CANON = {'Hospital & Surgical': ['GHS'], 'Major Medical': ['EMM'], 'Outpatient GP': ['GOPC'],
         'Outpatient Specialist': ['GOSC'], 'Dental': ['GD'], 'Term Life': ['GTL'],
         'Critical Illness': ['GCI'], 'Personal Accident': ['GPA']}

# V1.2 benefit ids -> canonical lines. A combined V1.2 line fills each canonical line it covers.
BENEFIT = {
    'ci.conditions': ['GCI_CONDITIONS'], 'ci.sum_assured': ['GCI_SUM_ASSURED'],
    'dental.annual_limit': ['GD_ANNUAL_LIMIT'], 'dental.copay': ['GD_CO_PAYMENT'],
    'ghs.annual_limit': ['GHS_ANNUAL_LIMIT'], 'ghs.cancer_kidney': ['GHS_OP_CANCER', 'GHS_OP_KIDNEY'],
    'ghs.copay': ['GHS_CO_PAYMENT'], 'ghs.death': ['GHS_DEATH'], 'ghs.emergency_op': ['GHS_EMERG_ACC_OP'],
    'ghs.hospital_type': ['GHS_HOSPITAL_TYPE'], 'ghs.icu': ['GHS_ICU'],
    'ghs.inpatient': ['GHS_HOSP_MISC', 'GHS_SURGICAL', 'GHS_INHOSP_DOCTOR'],
    'ghs.overseas': ['GHS_OVERSEAS'], 'ghs.pre_post': ['GHS_PRE_HOSP', 'GHS_POST_HOSP'],
    'ghs.pregnancy': ['GHS_MISCARRIAGE'], 'ghs.psychiatric': ['GHS_PSYCH'], 'ghs.rehab': ['GHS_REHAB'],
    'ghs.room_board': ['GHS_ROOM_BOARD'], 'ghs.surgical_implants': ['GHS_IMPLANTS'],
    'gmm.annual_limit': ['EMM_ANNUAL_LIMIT'], 'gmm.copay': ['EMM_CO_INSURANCE'],
    'gp.ae': ['GOPC_AE_VISIT'], 'gp.copay': ['GOPC_CO_PAYMENT'], 'gp.nonpanel': ['GOPC_NONPANEL'],
    'gp.panel': ['GOPC_PANEL_VISIT'], 'gp.tcm': ['GOPC_TCM'], 'gtl.sum_assured': ['GTL_SUM_ASSURED'],
    'pa.death': ['GPA_AD'], 'pa.medical': ['GPA_MEDICAL_EXP'], 'pa.tpd': ['GPA_PTD'],
    'sp.annual_limit': ['GOSC_ANNUAL_LIMIT'], 'sp.copay': ['GOSC_CO_PAYMENT'],
    'sp.diagnostic': ['GOSC_DIAGNOSTIC'], 'sp.physio': ['GOSC_PHYSIO'],
    # ghs.limit_basis says how a limit is counted, not a benefit; it goes into the notes.
}
NUM = re.compile(r'^\s*\$?\s*([\d,]+(?:\.\d+)?)\s*$')


def call(method, path, body=None, prefer='return=minimal'):
    a = ['curl', '-s', '-X', method, f'{SB}/{path}', '-H', f'apikey: {K}', '-H', f'Authorization: Bearer {K}',
         '-H', 'Content-Type: application/json', '-H', f'Prefer: {prefer}', '-w', '\n%{http_code}']
    if body is not None:
        with open('/tmp/_v12.json', 'w') as f: json.dump(body, f)
        a += ['--data-binary', '@/tmp/_v12.json']
    out = subprocess.run(a, capture_output=True, text=True).stdout
    txt, _, code = out.rpartition('\n')
    if not code.startswith('2'): sys.exit(f'{method} {path} -> {code}: {txt[:300]}')
    return json.loads(txt) if txt.strip() else None


def norm(s): return re.sub(r'[^a-z0-9]+', '', re.sub(r'\b(insurance|assurance|singapore|pte|ltd|limited|company|co|general|group|the)\b', ' ', (s or '').lower()))


def read_v12():
    wb = load_workbook(V12, data_only=True)
    def sheet(name, cols):
        rows = [[c.value for c in r] for r in wb[name].iter_rows(max_col=cols)]
        return [dict(zip(rows[0], r)) for r in rows[1:] if r[0]]
    return sheet('Rates', 11), sheet('BenefitData', 9)


def main():
    rates, benefits = read_v12()
    raw = open(V12, 'rb').read()
    v12_src = {'file': os.path.basename(V12), 'sha256': hashlib.sha256(raw).hexdigest(), 'md5': hashlib.md5(raw).hexdigest()}
    drive = {f['filename']: f for f in json.load(open('/tmp/drive_sources.json'))} if os.path.exists('/tmp/drive_sources.json') else {}
    companies = call('GET', 'companies?kind=eq.insurer&select=id,company_name', prefer='return=representation')
    keys = call('GET', 'insurers?select=id,name', prefer='return=representation')

    for v12_name, spec in INSURERS.items():
        R = [r for r in rates if r['Insurer'] == v12_name]
        B = [b for b in benefits if b['Insurer'] == v12_name]
        company = next((c for c in companies if c['company_name'] == spec['company']), None)
        if not company: sys.exit(f"{spec['company']} is not in Companies -> Insurers")
        key = next((k['id'] for k in keys if norm(k['name']) == norm(company['company_name'])), None)
        if key is None and APPLY:
            # The same key row src/lib/insurers.ts would create on first use.
            key = call('POST', 'insurers?on_conflict=name', {'name': spec['display'], 'status': 'active'},
                       prefer='resolution=merge-duplicates,return=representation')[0]['id']

        rate_rows = []
        for r in R:
            cls = re.search(r'\d', str(r['Class'])) if r['Class'] not in (None, '—') else None
            rate_rows.append({
                'product_code': r['Coverage'], 'plan_code': str(r['Plan option']),
                'member_type': {'EE': 'employee', 'Dep': 'dependant'}.get(str(r.get('Life'))),
                'band_label': f"{r['AgeFrom']} to {r['AgeTo']}" if r['AgeFrom'] is not None else 'All ages',
                'age_min': r['AgeFrom'] if r['AgeFrom'] is not None else 0, 'age_max': r['AgeTo'],
                'premium': round(float(r['Net premium']), 6), 'renewal_only': str(r['Basis']) == 'Renewal',
                'dimensions': {'occupation_class': cls.group(0)} if cls else {},
                'canon_codes': CANON[r['Coverage']],
            })
        plans = {}
        for r in R: plans.setdefault((r['Coverage'], str(r['Plan option'])), {})
        for b in B:
            p = plans.get((b['Coverage'], str(b['Plan'])))
            if p is None: continue
            if b['Benefit_ID'] == 'ghs.hospital_type': p['hospital_type'] = str(b['Value'])
            if b['Benefit_ID'] == 'ghs.room_board': p['beds'] = str(b['Value'])
            if b['Benefit_ID'] in ('ghs.copay',): p['co_payment'] = str(b['Value'])
        ben_rows, notes = [], []
        for b in B:
            if b['Benefit_ID'] == 'ghs.limit_basis':
                notes.append(f"{b['Plan']}: limit basis {b['Value']}"); continue
            for code in BENEFIT.get(b['Benefit_ID'], []):
                m = NUM.match(str(b['Value']))
                ben_rows.append({'product_code': b['Coverage'], 'plan_code': str(b['Plan']), 'category': CANON[b['Coverage']][0],
                                 'benefit_name': b['Benefit_Label'], 'value_text': str(b['Value']),
                                 'value_numeric': float(m.group(1).replace(',', '')) if m else None,
                                 'notes': b.get('Note') or None, 'canon_benefit': code, 'canon_codes': CANON[b['Coverage']]})

        products = list(dict.fromkeys(r['product_code'] for r in rate_rows))
        sources = []
        for fn, kind in ((v12_src['file'], 'workbook'), (spec['brochure'], 'brochure')):
            if fn and fn in drive:
                f = drive[fn]
                sources.append({'driveFileId': f['driveFileId'], 'filename': fn, 'insurer': f['insurer'], 'planYear': f['planYear'],
                                'kind': kind, 'md5': f['md5'], 'readAt': datetime.datetime.now(datetime.timezone.utc).isoformat()})
        rules = {'age_basis': spec['age_basis'].replace('_', ' '),
                 'gst_treatment': {'treatment': 'exclusive', 'conversion_factor': None, 'read_from': spec['gst_note']},
                 'age_basis_source': spec['age_note'], 'verification': spec['verification'],
                 'workbook': v12_src, 'sources': sources, 'notes': notes}

        existing = call('GET', f"gb_rate_tables?insurer_name=eq.{spec['display'].replace(' ', '%20')}&select=id", prefer='return=representation')
        print(f"\n{v12_name} -> {spec['display']} ({'replace ' + existing[0]['id'][:8] if existing else 'new table'})")
        print(f"  rates {len(rate_rows)} | plans {len(plans)} | benefit lines {len(ben_rows)} | products: {', '.join(products)}")
        print(f"  age basis {spec['age_basis']} | verification {spec['verification']['status']} | sources: {[s['filename'] for s in sources]}")
        if not APPLY: continue

        meta = {'insurer_id': key, 'insurer_name': spec['display'], 'product_code': ' · '.join(products),
                'product_name': spec['product_name'], 'plan_year': spec['plan_year'], 'age_basis': spec['age_basis'],
                'status': 'approved', 'source_pdf_name': spec['brochure'], 'rules': rules, 'rules_status': 'approved',
                'rules_updated_at': datetime.datetime.now(datetime.timezone.utc).isoformat(),
                'approved_at': datetime.datetime.now(datetime.timezone.utc).isoformat()}
        if existing:
            tid = existing[0]['id']
            call('PATCH', f'gb_rate_tables?id=eq.{tid}', meta)
            for t in ('gb_rates', 'gb_plans', 'gb_benefits'): call('DELETE', f'{t}?rate_table_id=eq.{tid}')
        else:
            tid = call('POST', 'gb_rate_tables', meta, prefer='return=representation')[0]['id']
        for i in range(0, len(rate_rows), 500):
            call('POST', 'gb_rates', [{**r, 'rate_table_id': tid, 'insurer_id': key, 'insurer_name': spec['display']} for r in rate_rows[i:i + 500]])
        call('POST', 'gb_plans', [{'rate_table_id': tid, 'product_code': c, 'plan_code': p, 'plan_name': p,
                                   'hospital_type': a.get('hospital_type'), 'beds': a.get('beds'), 'co_payment': a.get('co_payment'),
                                   'canon_codes': CANON[c], 'sort_order': i} for i, ((c, p), a) in enumerate(plans.items())])
        if ben_rows: call('POST', 'gb_benefits', [{**b, 'rate_table_id': tid} for b in ben_rows])
        print(f'  APPLIED -> rate table {tid}')
    if not APPLY: print('\nDRY RUN — nothing written. Re-run with --apply.')


if __name__ == '__main__':
    main()
